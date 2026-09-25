import io
import re
from datetime import date

import pandas as pd
import streamlit as st

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Membership & Orientation System",
    page_icon="📋",
    layout="wide"
)

st.title("📋 Membership & Orientation Tracking System")

st.write(
    "Upload the forms/master list and attendance files. "
    "The system will compare names, check duplicates, "
    "segregate records by status, and generate reports."
)


# ============================================================
# NAME NORMALIZATION
# ============================================================

def normalize_name(value):
    """
    Normalize names so that differences in capitalization,
    punctuation, commas, hyphens, and extra spaces do not
    prevent matching.
    """

    if value is None or pd.isna(value):
        return ""

    value = str(value).upper().strip()

    # Remove common punctuation
    value = re.sub(
        r"[,.;:/\\\-]+",
        " ",
        value
    )

    # Remove apostrophes
    value = value.replace("'", "")

    # Normalize multiple spaces
    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


def name_tokens(value):

    value = normalize_name(value)

    if not value:
        return []

    return value.split()


def normalize_part(value):

    return normalize_name(value)


# ============================================================
# STANDARDIZE COLUMN NAMES
# ============================================================

def standardize_columns(df):

    rename_map = {}

    for col in df.columns:

        original = str(col).strip()

        key = re.sub(
            r"[^A-Z0-9]",
            "",
            original.upper()
        )

        # ----------------------------------------------------
        # FORM / MASTER LIST
        # ----------------------------------------------------

        if key in [
            "LASTNAME",
            "SURNAME",
            "FAMILYNAME"
        ]:
            rename_map[col] = "Last Name"

        elif key in [
            "FIRSTNAME",
            "GIVENNAME",
            "GIVEN"
        ]:
            rename_map[col] = "First Name"

        elif key in [
            "MIDDLENAME",
            "MIDDLE",
            "MIDDLEINITIAL",
            "MI"
        ]:
            rename_map[col] = "Middle Name"

        elif key in [
            "NICKNAME",
            "NICK"
        ]:
            rename_map[col] = "Nickname"

        elif key in [
            "DATEOFBIRTH",
            "DOB",
            "BIRTHDATE"
        ]:
            rename_map[col] = "Date of Birth"

        elif key == "AGE":
            rename_map[col] = "Age"

        elif key == "CIVILSTATUS":
            rename_map[col] = "Civil Status"

        elif key in [
            "GENDER",
            "SEX"
        ]:
            rename_map[col] = "Gender"

        elif key in [
            "ADDRESS",
            "HOMEADDRESS"
        ]:
            rename_map[col] = "Address"

        elif key in [
            "CONTACTNO",
            "CONTACTNUMBER",
            "PHONE",
            "PHONENUMBER",
            "MOBILENUMBER"
        ]:
            rename_map[col] = "Contact No."

        elif key in [
            "FBACCOUNT",
            "FACEBOOK",
            "FACEBOOKACCOUNT"
        ]:
            rename_map[col] = "FB Account"

        elif key in [
            "FATHERSNAME",
            "FATHERNAME"
        ]:
            rename_map[col] = "Fathers Name"

        elif key in [
            "MOTHERSNAME",
            "MOTHERNAME"
        ]:
            rename_map[col] = "Mothers Name"

        elif key == "RELIGION":
            rename_map[col] = "Religion"

        elif key in [
            "NOOFCHILDREN",
            "NUMBEROFCHILDREN",
            "CHILDREN"
        ]:
            rename_map[col] = "No of Children"

        elif key in [
            "SPOUSENAME",
            "SPOUSE"
        ]:
            rename_map[col] = "Spouse Name"

        # ----------------------------------------------------
        # ATTENDANCE
        # ----------------------------------------------------

        elif key in [
            "FULLNAME",
            "NAME",
            "COMPLETENAME",
            "ATTENDEENAME",
            "ATTENDEE"
        ]:
            rename_map[col] = "Full Name"

        elif key in [
            "ATTENDANCE",
            "ATTENDED",
            "PRESENT"
        ]:
            rename_map[col] = "Attendance"

    return df.rename(
        columns=rename_map
    )


# ============================================================
# READ FILE
# ============================================================

def read_file(uploaded_file):

    filename = uploaded_file.name.lower()

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    if filename.endswith(".csv"):

        df = pd.read_csv(
            uploaded_file,
            dtype=str
        )

        return standardize_columns(df)

    # --------------------------------------------------------
    # EXCEL
    # --------------------------------------------------------

    excel = pd.ExcelFile(
        uploaded_file
    )

    all_sheets = []

    for sheet in excel.sheet_names:

        df = pd.read_excel(
            uploaded_file,
            sheet_name=sheet,
            dtype=str
        )

        if not df.empty:

            df = standardize_columns(
                df
            )

            all_sheets.append(df)

    if not all_sheets:

        return pd.DataFrame()

    return pd.concat(
        all_sheets,
        ignore_index=True
    )


# ============================================================
# PREPARE FORMS / MASTER LIST
# ============================================================

def prepare_master(df):

    df = standardize_columns(
        df.copy()
    )

    required_columns = [
        "Last Name",
        "First Name",
        "Middle Name"
    ]

    for col in required_columns:

        if col not in df.columns:

            df[col] = ""

    return df


# ============================================================
# PREPARE ATTENDANCE
# ============================================================

def prepare_attendance(df):

    df = standardize_columns(
        df.copy()
    )

    # Attendance should normally contain:
    #
    # Full Name
    #
    # Nothing else is required.

    if "Full Name" not in df.columns:

        # Try to identify a likely name column
        possible_columns = []

        for col in df.columns:

            key = re.sub(
                r"[^A-Z0-9]",
                "",
                str(col).upper()
            )

            if key not in [
                "NO",
                "NUMBER",
                "ID",
                "DATE",
                "ATTENDANCE",
                "ATTENDED",
                "PRESENT"
            ]:

                possible_columns.append(col)

        if possible_columns:

            df = df.rename(
                columns={
                    possible_columns[0]:
                    "Full Name"
                }
            )

    if "Full Name" not in df.columns:

        raise ValueError(
            "Attendance file must contain "
            "a Full Name or Name column."
        )

    return df


# ============================================================
# DISPLAY NAME
# ============================================================

def display_master_name(row):

    first = str(
        row.get(
            "First Name",
            ""
        )
    ).strip()

    middle = str(
        row.get(
            "Middle Name",
            ""
        )
    ).strip()

    last = str(
        row.get(
            "Last Name",
            ""
        )
    ).strip()

    parts = []

    if first and first.lower() != "nan":
        parts.append(first)

    if middle and middle.lower() != "nan":
        parts.append(middle)

    if last and last.lower() != "nan":
        parts.append(last)

    return " ".join(parts)


# ============================================================
# NAME VARIATIONS
# ============================================================

def generate_name_variations(row):

    last = normalize_part(
        row.get(
            "Last Name",
            ""
        )
    )

    first = normalize_part(
        row.get(
            "First Name",
            ""
        )
    )

    middle = normalize_part(
        row.get(
            "Middle Name",
            ""
        )
    )

    variations = set()

    # --------------------------------------------------------
    # FIRST + LAST
    # --------------------------------------------------------

    if first and last:

        variations.add(
            normalize_name(
                f"{first} {last}"
            )
        )

        variations.add(
            normalize_name(
                f"{last} {first}"
            )
        )

    # --------------------------------------------------------
    # FIRST + MIDDLE + LAST
    # --------------------------------------------------------

    if first and middle and last:

        variations.add(
            normalize_name(
                f"{first} {middle} {last}"
            )
        )

        variations.add(
            normalize_name(
                f"{last} {first} {middle}"
            )
        )

        variations.add(
            normalize_name(
                f"{last} {middle} {first}"
            )
        )

        variations.add(
            normalize_name(
                f"{first} {last} {middle}"
            )
        )

        variations.add(
            normalize_name(
                f"{middle} {first} {last}"
            )
        )

    # --------------------------------------------------------
    # MIDDLE INITIAL
    # --------------------------------------------------------

    if first and middle and last:

        middle_words = middle.split()

        if middle_words:

            initial = middle_words[0][0]

            variations.add(
                normalize_name(
                    f"{first} {initial} {last}"
                )
            )

            variations.add(
                normalize_name(
                    f"{last} {first} {initial}"
                )
            )

            variations.add(
                normalize_name(
                    f"{last} {initial} {first}"
                )
            )

            variations.add(
                normalize_name(
                    f"{first} {last} {initial}"
                )
            )

    return {
        v for v in variations
        if v
    }


# ============================================================
# EXACT NAME MATCH
# ============================================================

def exact_name_match(
    master_row,
    attendance_name
):

    attendance_normalized = normalize_name(
        attendance_name
    )

    if not attendance_normalized:

        return False

    variations = generate_name_variations(
        master_row
    )

    return (
        attendance_normalized
        in variations
    )


# ============================================================
# TOKEN MATCH
# ============================================================

def token_match(
    master_row,
    attendance_name
):

    attendance = normalize_name(
        attendance_name
    )

    attendance_tokens = set(
        name_tokens(attendance)
    )

    if not attendance_tokens:

        return False

    first = normalize_part(
        master_row.get(
            "First Name",
            ""
        )
    )

    last = normalize_part(
        master_row.get(
            "Last Name",
            ""
        )
    )

    if not first or not last:

        return False

    first_parts = first.split()
    last_parts = last.split()

    # Every first-name token must exist
    first_ok = all(
        part in attendance_tokens
        for part in first_parts
    )

    # Every surname token must exist
    last_ok = all(
        part in attendance_tokens
        for part in last_parts
    )

    if not first_ok or not last_ok:

        return False

    return True


# ============================================================
# FIND MATCHES
# ============================================================

def find_matching_attendance(
    master_row,
    attendance
):

    exact_matches = []
    token_matches = []

    for idx, attendance_row in (
        attendance.iterrows()
    ):

        attendance_name = (
            attendance_row.get(
                "Full Name",
                ""
            )
        )

        # Exact variation match
        if exact_name_match(
            master_row,
            attendance_name
        ):

            exact_matches.append(
                (
                    idx,
                    attendance_row
                )
            )

        # Token match
        elif token_match(
            master_row,
            attendance_name
        ):

            token_matches.append(
                (
                    idx,
                    attendance_row
                )
            )

    if exact_matches:

        return (
            exact_matches,
            "EXACT"
        )

    if token_matches:

        return (
            token_matches,
            "NAME PARTS"
        )

    return (
        [],
        "NONE"
    )


# ============================================================
# FIND DUPLICATES IN FORMS
# ============================================================

def find_duplicates(master):

    if master.empty:

        return pd.DataFrame()

    df = master.copy()

    for col in [
        "Last Name",
        "First Name",
        "Middle Name"
    ]:

        if col not in df.columns:

            df[col] = ""

    df["_last"] = (
        df["Last Name"]
        .apply(normalize_part)
    )

    df["_first"] = (
        df["First Name"]
        .apply(normalize_part)
    )

    df["_middle"] = (
        df["Middle Name"]
        .apply(normalize_part)
    )

    df["_name"] = (
        df["_last"]
        + "|"
        + df["_first"]
        + "|"
        + df["_middle"]
    )

    duplicate_mask = (
        df["_name"].ne("||")
        &
        df.duplicated(
            "_name",
            keep=False
        )
    )

    duplicates = df[
        duplicate_mask
    ].copy()

    if duplicates.empty:

        return pd.DataFrame()

    duplicates["Duplicate Type"] = (
        "Duplicate in Forms"
    )

    return remove_internal_columns(
        duplicates
    )


# ============================================================
# REMOVE INTERNAL COLUMNS
# ============================================================

def remove_internal_columns(df):

    if df.empty:

        return df

    return df[
        [
            col
            for col in df.columns
            if not str(col).startswith("_")
        ]
    ]


# ============================================================
# PROCESS EVERYTHING
# ============================================================

def process_records(
    master,
    attendance
):

    master = prepare_master(
        master
    )

    attendance = prepare_attendance(
        attendance
    )

    # --------------------------------------------------------
    # RESULT CONTAINERS
    # --------------------------------------------------------

    oriented_with_form = []

    for_orientation = []

    oriented_without_form = []

    verification = []

    # --------------------------------------------------------
    # TRACK USED RECORDS
    # --------------------------------------------------------

    matched_attendance = set()

    matched_master = set()

    # ========================================================
    # PASS 1
    #
    # FORMS -> ATTENDANCE
    #
    # Every person who submitted a form is checked against
    # attendance.
    #
    # Match = ORIENTED WITH FORM
    # No match = FOR ORIENTATION
    # ========================================================

    for master_idx, master_row in (
        master.iterrows()
    ):

        matches, match_type = (
            find_matching_attendance(
                master_row,
                attendance
            )
        )

        # ----------------------------------------------------
        # FORM PERSON NOT FOUND IN ATTENDANCE
        # ----------------------------------------------------

        if not matches:

            record = master_row.copy()

            record["Status"] = (
                "For Orientation"
            )

            for_orientation.append(
                record
            )

            continue

        # ----------------------------------------------------
        # MULTIPLE ATTENDANCE MATCHES
        # ----------------------------------------------------

        if len(matches) > 1:

            record = master_row.copy()

            record["Status"] = (
                "For Verification"
            )

            record[
                "Possible Attendance Matches"
            ] = " | ".join(
                str(
                    row.get(
                        "Full Name",
                        ""
                    )
                )
                for _, row in matches
            )

            record[
                "Verification Reason"
            ] = (
                "One form matches multiple "
                "attendance records."
            )

            verification.append(
                record
            )

            continue

        # ----------------------------------------------------
        # ONE ATTENDANCE MATCH
        # ----------------------------------------------------

        attendance_idx, attendance_row = (
            matches[0]
        )

        # Attendance already assigned
        if attendance_idx in matched_attendance:

            record = master_row.copy()

            record["Status"] = (
                "For Verification"
            )

            record[
                "Attendance Name"
            ] = attendance_row.get(
                "Full Name",
                ""
            )

            record[
                "Verification Reason"
            ] = (
                "Attendance record was already "
                "matched to another form."
            )

            verification.append(
                record
            )

            continue

        # Mark both as matched
        matched_attendance.add(
            attendance_idx
        )

        matched_master.add(
            master_idx
        )

        # ----------------------------------------------------
        # ORIENTED WITH FORM
        # ----------------------------------------------------

        record = master_row.copy()

        record["Status"] = (
            "Oriented With Form"
        )

        record["Match Method"] = (
            match_type
        )

        record["Attendance Name"] = (
            attendance_row.get(
                "Full Name",
                ""
            )
        )

        oriented_with_form.append(
            record
        )

    # ========================================================
    # PASS 2
    #
    # ATTENDANCE -> FORMS
    #
    # Attendance records not matched in Pass 1 are checked
    # against the forms again.
    #
    # If no form exists:
    #
    # ORIENTED WITHOUT FORM
    #
    # If multiple forms match:
    #
    # FOR VERIFICATION
    # ========================================================

    for attendance_idx, attendance_row in (
        attendance.iterrows()
    ):

        # Already matched to a form
        if attendance_idx in matched_attendance:

            continue

        attendance_name = (
            attendance_row.get(
                "Full Name",
                ""
            )
        )

        if not normalize_name(
            attendance_name
        ):

            continue

        possible_master = []

        # ----------------------------------------------------
        # Search forms
        # ----------------------------------------------------

        for master_idx, master_row in (
            master.iterrows()
        ):

            if exact_name_match(
                master_row,
                attendance_name
            ):

                possible_master.append(
                    (
                        master_idx,
                        master_row,
                        "EXACT"
                    )
                )

            elif token_match(
                master_row,
                attendance_name
            ):

                possible_master.append(
                    (
                        master_idx,
                        master_row,
                        "NAME PARTS"
                    )
                )

        # ----------------------------------------------------
        # NO FORM
        #
        # THIS IS ORIENTED WITHOUT FORM
        # ----------------------------------------------------

        if not possible_master:

            record = attendance_row.copy()

            record["Status"] = (
                "Oriented Without Form"
            )

            record[
                "Attendance Name"
            ] = attendance_name

            oriented_without_form.append(
                record
            )

            matched_attendance.add(
                attendance_idx
            )

        # ----------------------------------------------------
        # ONE FORM MATCH
        # ----------------------------------------------------

        elif len(possible_master) == 1:

            master_idx, master_row, method = (
                possible_master[0]
            )

            # If the form was already matched,
            # this attendance record is a duplicate.
            if master_idx in matched_master:

                record = attendance_row.copy()

                record["Status"] = (
                    "For Verification"
                )

                record[
                    "Possible Master Record"
                ] = display_master_name(
                    master_row
                )

                record[
                    "Verification Reason"
                ] = (
                    "This form already has an "
                    "attendance match."
                )

                verification.append(
                    record
                )

                matched_attendance.add(
                    attendance_idx
                )

            else:

                # This can happen because Pass 1 may have
                # selected a different matching route.

                record = master_row.copy()

                record["Status"] = (
                    "Oriented With Form"
                )

                record["Match Method"] = (
                    method
                )

                record["Attendance Name"] = (
                    attendance_name
                )

                oriented_with_form.append(
                    record
                )

                matched_master.add(
                    master_idx
                )

                matched_attendance.add(
                    attendance_idx
                )

        # ----------------------------------------------------
        # MULTIPLE FORMS MATCH
        # ----------------------------------------------------

        else:

            record = attendance_row.copy()

            record["Status"] = (
                "For Verification"
            )

            record[
                "Possible Master Records"
            ] = " | ".join(
                display_master_name(
                    master_row
                )
                for _, master_row, _ in (
                    possible_master
                )
            )

            record[
                "Verification Reason"
            ] = (
                "Attendance name matches "
                "multiple forms."
            )

            verification.append(
                record
            )

            matched_attendance.add(
                attendance_idx
            )

    return (
        pd.DataFrame(
            oriented_with_form
        ),
        pd.DataFrame(
            for_orientation
        ),
        pd.DataFrame(
            oriented_without_form
        ),
        pd.DataFrame(
            verification
        )
    )


# ============================================================
# SUMMARY
# ============================================================

def create_summary(
    oriented_with_form,
    for_orientation,
    oriented_without_form,
    duplicates,
    verification
):

    return pd.DataFrame({

        "Category": [
            "Oriented With Form",
            "For Orientation",
            "Oriented Without Form",
            "Duplicates",
            "For Verification"
        ],

        "Count": [
            len(oriented_with_form),
            len(for_orientation),
            len(oriented_without_form),
            len(duplicates),
            len(verification)
        ]
    })


# ============================================================
# CREATE EXCEL REPORT
# ============================================================

def create_excel(
    summary,
    oriented_with_form,
    for_orientation,
    oriented_without_form,
    duplicates,
    verification
):

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        oriented_with_form.to_excel(
            writer,
            sheet_name="Oriented With Form",
            index=False
        )

        for_orientation.to_excel(
            writer,
            sheet_name="For Orientation",
            index=False
        )

        oriented_without_form.to_excel(
            writer,
            sheet_name="Oriented Without Form",
            index=False
        )

        duplicates.to_excel(
            writer,
            sheet_name="Duplicates",
            index=False
        )

        verification.to_excel(
            writer,
            sheet_name="For Verification",
            index=False
        )

    output.seek(0)

    return output


# ============================================================
# POWERPOINT TITLE
# ============================================================

def add_slide_title(
    slide,
    title,
    subtitle=""
):

    box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(0.35),
        Inches(12),
        Inches(0.7)
    )

    tf = box.text_frame

    tf.clear()

    p = tf.paragraphs[0]

    p.text = title

    p.font.size = Pt(28)
    p.font.bold = True

    p.font.color.rgb = RGBColor(
        31,
        78,
        121
    )

    if subtitle:

        sub = slide.shapes.add_textbox(
            Inches(0.6),
            Inches(1.0),
            Inches(12),
            Inches(0.4)
        )

        tf2 = sub.text_frame

        tf2.clear()

        p2 = tf2.paragraphs[0]

        p2.text = subtitle

        p2.font.size = Pt(14)

        p2.font.color.rgb = RGBColor(
            100,
            100,
            100
        )


# ============================================================
# POWERPOINT REPORT
# ============================================================

def create_powerpoint(
    barangay,
    encoded_date,
    summary
):

    prs = Presentation()

    # ========================================================
    # TITLE SLIDE
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Membership & Orientation Report",
        f"Barangay {barangay}, Alaminos"
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(11),
        Inches(3)
    )

    tf = box.text_frame

    tf.clear()

    lines = [
        f"Form Encoding Date: {encoded_date}",
        "",
        "Membership and Orientation Reconciliation"
    ]

    for i, line in enumerate(lines):

        p = (
            tf.paragraphs[0]
            if i == 0
            else tf.add_paragraph()
        )

        p.text = line

        p.alignment = PP_ALIGN.CENTER

        p.font.size = Pt(
            20 if i == 0 else 16
        )

        if i == 0:

            p.font.bold = True

    # ========================================================
    # SUMMARY SLIDE
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Overall Summary"
    )

    colors = {

        "Oriented With Form":
            RGBColor(
                46,
                125,
                50
            ),

        "For Orientation":
            RGBColor(
                255,
                152,
                0
            ),

        "Oriented Without Form":
            RGBColor(
                33,
                150,
                243
            ),

        "Duplicates":
            RGBColor(
                198,
                40,
                40
            ),

        "For Verification":
            RGBColor(
                123,
                31,
                162
            )
    }

    x_positions = [
        0.15,
        2.65,
        5.15,
        7.65,
        10.15
    ]

    for i, row in summary.iterrows():

        category = row[
            "Category"
        ]

        box = slide.shapes.add_textbox(
            Inches(
                x_positions[i]
            ),
            Inches(1.9),
            Inches(2.35),
            Inches(2)
        )

        tf = box.text_frame

        tf.clear()

        p = tf.paragraphs[0]

        p.text = str(
            int(row["Count"])
        )

        p.alignment = PP_ALIGN.CENTER

        p.font.size = Pt(34)

        p.font.bold = True

        p.font.color.rgb = colors.get(
            category,
            RGBColor(
                50,
                50,
                50
            )
        )

        p2 = tf.add_paragraph()

        p2.text = category

        p2.alignment = PP_ALIGN.CENTER

        p2.font.size = Pt(11)

        p2.font.bold = True

    # ========================================================
    # CHART
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Orientation Status"
    )

    chart_data = CategoryChartData()

    chart_data.categories = [
        "Oriented With Form",
        "For Orientation",
        "Without Form"
    ]

    values = []

    for category in [
        "Oriented With Form",
        "For Orientation",
        "Oriented Without Form"
    ]:

        found = summary[
            summary["Category"]
            == category
        ]

        if found.empty:

            values.append(0)

        else:

            values.append(
                int(
                    found[
                        "Count"
                    ].iloc[0]
                )
            )

    chart_data.add_series(
        "People",
        values
    )

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1.5),
        Inches(11),
        Inches(5),
        chart_data
    )

    chart = chart_shape.chart

    chart.has_legend = False

    # ========================================================
    # FOLLOW-UP SLIDE
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Follow-up Required"
    )

    row = summary[
        summary["Category"]
        == "For Orientation"
    ]

    count = (
        int(
            row["Count"].iloc[0]
        )
        if not row.empty
        else 0
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.8),
        Inches(11),
        Inches(3)
    )

    tf = box.text_frame

    tf.clear()

    p = tf.paragraphs[0]

    p.text = str(count)

    p.alignment = PP_ALIGN.CENTER

    p.font.size = Pt(60)

    p.font.bold = True

    p.font.color.rgb = RGBColor(
        255,
        152,
        0
    )

    p2 = tf.add_paragraph()

    p2.text = (
        "People with forms who have "
        "not yet attended orientation"
    )

    p2.alignment = PP_ALIGN.CENTER

    p2.font.size = Pt(18)

    # ========================================================
    # WITHOUT FORM SLIDE
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Oriented Without Form"
    )

    row = summary[
        summary["Category"]
        == "Oriented Without Form"
    ]

    count = (
        int(
            row["Count"].iloc[0]
        )
        if not row.empty
        else 0
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.8),
        Inches(11),
        Inches(3)
    )

    tf = box.text_frame

    tf.clear()

    p = tf.paragraphs[0]

    p.text = str(count)

    p.alignment = PP_ALIGN.CENTER

    p.font.size = Pt(60)

    p.font.bold = True

    p.font.color.rgb = RGBColor(
        33,
        150,
        243
    )

    p2 = tf.add_paragraph()

    p2.text = (
        "People found in attendance "
        "but with no matching form"
    )

    p2.alignment = PP_ALIGN.CENTER

    p2.font.size = Pt(18)

    # ========================================================
    # DATA QUALITY SLIDE
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Data Quality"
    )

    duplicate_row = summary[
        summary["Category"]
        == "Duplicates"
    ]

    verification_row = summary[
        summary["Category"]
        == "For Verification"
    ]

    duplicate_count = (
        int(
            duplicate_row[
                "Count"
            ].iloc[0]
        )
        if not duplicate_row.empty
        else 0
    )

    verification_count = (
        int(
            verification_row[
                "Count"
            ].iloc[0]
        )
        if not verification_row.empty
        else 0
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.8),
        Inches(11),
        Inches(4)
    )

    tf = box.text_frame

    tf.clear()

    items = [

        f"Duplicate records: "
        f"{duplicate_count}",

        f"Records for verification: "
        f"{verification_count}",

        "",

        "Records marked for verification "
        "should be reviewed before finalizing "
        "the report."
    ]

    for i, item in enumerate(items):

        p = (
            tf.paragraphs[0]
            if i == 0
            else tf.add_paragraph()
        )

        p.text = item

        p.font.size = Pt(
            20 if i < 2 else 16
        )

        p.space_after = Pt(10)

    # ========================================================
    # NEXT ACTIONS
    # ========================================================

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_slide_title(
        slide,
        "Next Actions"
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(11),
        Inches(4.8)
    )

    tf = box.text_frame

    tf.clear()

    actions = [

        "Follow up with people who have "
        "forms but have not attended orientation.",

        "Schedule pending individuals "
        "for orientation.",

        "Review records marked "
        "For Verification.",

        "Resolve duplicate records.",

        "Review people who attended "
        "without submitting a form.",

        "Update membership records "
        "after orientation."
    ]

    for i, action in enumerate(actions):

        p = (
            tf.paragraphs[0]
            if i == 0
            else tf.add_paragraph()
        )

        p.text = "• " + action

        p.font.size = Pt(18)

        p.space_after = Pt(12)

    # ========================================================
    # SAVE
    # ========================================================

    output = io.BytesIO()

    prs.save(output)

    output.seek(0)

    return output


# ============================================================
# INPUT SECTION
# ============================================================

st.divider()

st.header("1. Batch Information")

col1, col2 = st.columns(2)

with col1:

    barangay = st.text_input(
        "Barangay",
        placeholder="Example: Poblacion"
    )

with col2:

    encoded_date = st.date_input(
        "Form Encoding Date",
        value=date.today()
    )

st.caption(
    "This is the date the collected forms were encoded "
    "into the system."
)


# ============================================================
# FILE UPLOAD
# ============================================================

st.header("2. Upload Files")

master_file = st.file_uploader(
    "Forms / Master List",
    type=[
        "xlsx",
        "xls",
        "csv"
    ],
    help=(
        "This file contains the people who submitted "
        "forms. It should contain Last Name, First Name, "
        "and optionally Middle Name."
    )
)

attendance_files = st.file_uploader(
    "Orientation Attendance File(s)",
    type=[
        "xlsx",
        "xls",
        "csv"
    ],
    accept_multiple_files=True,
    help=(
        "Attendance files only need the person's Full Name. "
        "Middle name or middle initial may be missing."
    )
)


# ============================================================
# PROCESS
# ============================================================

if st.button(
    "🔄 Process & Generate Report",
    type="primary",
    use_container_width=True
):

    # --------------------------------------------------------
    # VALIDATION
    # --------------------------------------------------------

    if not barangay.strip():

        st.error(
            "Please enter the Barangay."
        )

        st.stop()

    if not master_file:

        st.error(
            "Please upload the Forms / Master List."
        )

        st.stop()

    if not attendance_files:

        st.error(
            "Please upload at least one "
            "attendance file."
        )

        st.stop()

    try:

        # ====================================================
        # READ FORMS
        # ====================================================

        with st.spinner(
            "Reading forms..."
        ):

            master = read_file(
                master_file
            )

            master = prepare_master(
                master
            )

        # ====================================================
        # READ ATTENDANCE
        # ====================================================

        attendance_parts = []

        with st.spinner(
            "Reading attendance files..."
        ):

            for file in attendance_files:

                df = read_file(
                    file
                )

                df = prepare_attendance(
                    df
                )

                # Keep track of source file
                df[
                    "Attendance Source"
                ] = file.name

                attendance_parts.append(
                    df
                )

        attendance = pd.concat(
            attendance_parts,
            ignore_index=True
        )

        # ====================================================
        # FIND DUPLICATES
        # ====================================================

        with st.spinner(
            "Checking duplicates..."
        ):

            duplicates = find_duplicates(
                master
            )

            # ------------------------------------------------
            # ATTENDANCE DUPLICATES
            # ------------------------------------------------

            if (
                not attendance.empty
                and
                "Full Name"
                in attendance.columns
            ):

                attendance_names = (
                    attendance[
                        "Full Name"
                    ]
                    .apply(
                        normalize_name
                    )
                )

                duplicate_mask = (
                    attendance_names.ne("")
                    &
                    attendance_names.duplicated(
                        keep=False
                    )
                )

                attendance_dupes = (
                    attendance[
                        duplicate_mask
                    ].copy()
                )

                if not attendance_dupes.empty:

                    attendance_dupes[
                        "Duplicate Type"
                    ] = (
                        "Duplicate in Attendance"
                    )

                    duplicates = pd.concat(
                        [
                            duplicates,
                            attendance_dupes
                        ],
                        ignore_index=True
                    )

        # ====================================================
        # PROCESS MATCHING
        # ====================================================

        with st.spinner(
            "Comparing forms and attendance..."
        ):

            (
                oriented_with_form,
                for_orientation,
                oriented_without_form,
                verification
            ) = process_records(
                master,
                attendance
            )

        # ====================================================
        # SUMMARY
        # ====================================================

        summary = create_summary(
            oriented_with_form,
            for_orientation,
            oriented_without_form,
            duplicates,
            verification
        )

        # ====================================================
        # SUCCESS
        # ====================================================

        st.success(
            "Processing complete!"
        )

        # ====================================================
        # SUMMARY METRICS
        # ====================================================

        st.header("3. Summary")

        cols = st.columns(5)

        for col, (_, row) in zip(
            cols,
            summary.iterrows()
        ):

            with col:

                st.metric(
                    row["Category"],
                    int(
                        row["Count"]
                    )
                )

        # ====================================================
        # SEGREGATED TABLES
        # ====================================================

        st.header(
            "4. Segregated Records"
        )

        tabs = st.tabs([
            "✅ Oriented With Form",
            "📌 For Orientation",
            "📝 Oriented Without Form",
            "⚠️ Duplicates",
            "🔎 For Verification"
        ])

        # ----------------------------------------------------
        # ORIENTED WITH FORM
        # ----------------------------------------------------

        with tabs[0]:

            st.dataframe(
                oriented_with_form,
                use_container_width=True,
                hide_index=True
            )

        # ----------------------------------------------------
        # FOR ORIENTATION
        # ----------------------------------------------------

        with tabs[1]:

            st.dataframe(
                for_orientation,
                use_container_width=True,
                hide_index=True
            )

        # ----------------------------------------------------
        # WITHOUT FORM
        # ----------------------------------------------------

        with tabs[2]:

            st.dataframe(
                oriented_without_form,
                use_container_width=True,
                hide_index=True
            )

        # ----------------------------------------------------
        # DUPLICATES
        # ----------------------------------------------------

        with tabs[3]:

            st.dataframe(
                duplicates,
                use_container_width=True,
                hide_index=True
            )

        # ----------------------------------------------------
        # VERIFICATION
        # ----------------------------------------------------

        with tabs[4]:

            st.dataframe(
                verification,
                use_container_width=True,
                hide_index=True
            )

        # ====================================================
        # EXCEL
        # ====================================================

        st.header(
            "5. Download Reports"
        )

        excel = create_excel(
            summary,
            oriented_with_form,
            for_orientation,
            oriented_without_form,
            duplicates,
            verification
        )

        st.download_button(
            "⬇️ Download Excel Report",
            data=excel,
            file_name=(
                f"{barangay}_"
                f"membership_report.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True
        )

        # ====================================================
        # POWERPOINT
        # ====================================================

        with st.spinner(
            "Creating PowerPoint presentation..."
        ):

            ppt = create_powerpoint(
                barangay,
                encoded_date.strftime(
                    "%B %d, %Y"
                ),
                summary
            )

        st.download_button(
            "📊 Download PowerPoint Report",
            data=ppt,
            file_name=(
                f"{barangay}_"
                f"orientation_report.pptx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            use_container_width=True
        )

    except Exception as e:

        st.error(
            "There was an error processing the files."
        )

        st.exception(e)
